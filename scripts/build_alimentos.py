"""Arma public/alimentos/{smae,renal}.json a partir del SMAE de Cecy y de scripts/europeos.csv.

Uso: python3 scripts/build_alimentos.py <Smae.xlsm> <SMAE_datos_a_revisar.xlsx> [quitados.xlsx]
"""
import csv, json, re, sys, unicodedata
from pathlib import Path
import openpyxl
from saludable import motivo

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "public" / "alimentos"

# Valor de referencia de 1 equivalente por subgrupo (kcal, proteína, lípidos, HC)
REF = {
    "Verduras": (25, 2, 0, 4), "Frutas": (60, 0, 0, 15),
    "Cereales S/G": (70, 2, 0, 15), "Cereales C/G": (115, 2, 5, 15),
    "Leguminosas": (120, 8, 1, 20),
    "A.O.A MBAG": (40, 7, 1, 0), "A.O.A BAG": (55, 7, 3, 0), "A.O.A MAG": (75, 7, 5, 0), "A.O.A AAG": (100, 7, 8, 0),
    "Leche descremada": (95, 9, 2, 12), "Leche semi": (110, 9, 4, 12), "Leche entera": (150, 9, 8, 12), "Leche C/azúcar": (200, 8, 5, 30),
    "Grasas S/Proteínas": (45, 0, 5, 0), "Grasas C/Proteínas": (70, 3, 5, 3),
    "Azucares S/grasa": (40, 0, 0, 10), "Azucares C/grasa": (85, 0, 5, 10),
}

# Comida chatarra que Cecy pidió omitir (botanas de bolsa, pan dulce empaquetado, donas, nuggets, refresco, dulces)
CHATARRA = re.compile(
    r"cheetos|doritos|sabritas|^fritos$|churritos|barcel|chicharr[oó]n de harina|frituras|gansito|marinela|bigote|pastelito|"
    r"cupcake|\bdonas?\b|nugget|papas fritas|papas sabritas|kkwatte|carlos v|o.reo|refresco(?! de dieta| en polvo sin)|"
    r"gomitas|caramelo|malvavisco|mentas|paleta de caramelo|pasitas con chocolate|sándwich quaker|nachos|empanizad|eclair|"
    r"galleta.*(chocolate|relleno cremoso|malvavisco|sandwich|sandwuich|wafers|cremosa)|marías gamesa cereale|panquecito|pay helado",
    re.I)

def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", s)).strip()

def slug(s):
    return norm(s).replace(" ", "-")[:48].strip("-")

FR = {"¼": .25, "⅓": 1/3, "½": .5, "⅔": 2/3, "¾": .75}

def num(v):
    if v is None: return None
    if isinstance(v, str) and v.strip() in FR: return FR[v.strip()]
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).replace(",", ".").strip())
    except ValueError: return None

def r1(v): return None if v is None else round(v, 1)

def frac(c):
    """0.333 -> 1/3 para que la paciente lo lea bien."""
    for v, t in ((1/4, "¼"), (1/3, "⅓"), (1/2, "½"), (2/3, "⅔"), (3/4, "¾")):
        w = int(c)
        if abs(c - w - v) < 0.02: return (f"{w}" if w else "") + t
    return f"{round(c, 2):g}"

ids = set()
def uid(base):
    i, k = base, 2
    while i in ids: i, k = f"{base}-{k}", k + 1
    ids.add(i); return i

def food(nombre, sub, cant, unidad, peso, k, p, l, h, f, origen, extra=None):
    d = {"id": uid(slug(nombre)), "n": re.sub(r"\s+", " ", str(nombre)).strip(), "g": sub, "c": round(cant, 3), "u": str(unidad).strip(),
         "pn": r1(peso), "k": r1(k), "p": r1(p), "l": r1(l), "h": r1(h), "o": origen}
    if f: d["f"] = r1(f)
    if extra: d.update({a: b for a, b in extra.items() if b is not None})
    return d

def smae(xlsm, revisar):
    wb = openpyxl.load_workbook(xlsm, read_only=True, data_only=True)
    quitar = set()
    for r in list(openpyxl.load_workbook(revisar, read_only=True).active.iter_rows(values_only=True))[1:]:
        if r[1]: quitar.add((r[0], norm(r[1]), round(num(r[2]) or 0, 1), str(r[3]).strip()))
    out, omit = [], {"revisar": 0, "aguacate california": 0, "sin datos": 0, "chatarra": [], "alcohol": 0}
    for r in list(wb["SMAE 5"].iter_rows(values_only=True))[1:]:
        nombre, sub = r[0], r[1]
        if not nombre or not sub: continue
        if sub == "Bebidas alcoholicas": omit["alcohol"] += 1; continue
        if (sub, norm(nombre), round(num(r[2]) or 0, 1), str(r[3]).strip()) in quitar: omit["revisar"] += 1; continue
        if "aguacate california" in norm(nombre): omit["aguacate california"] += 1; continue
        k, p, l, h, pn, c = (num(r[i]) for i in (6, 7, 8, 9, 5, 2))
        if sub != "Alimentos libres en energía" and (not k or c is None or not pn or (p, l, h) == (0, 0, 0)): omit["sin datos"] += 1; continue
        if sub == "Alimentos libres en energía" and (c is None or k is None): omit["sin datos"] += 1; continue
        if CHATARRA.search(nombre): omit["chatarra"].append(nombre); continue
        out.append(food(nombre, sub if sub != "Alimentos libres en energía" else "Libres", c, r[3] or "", pn, k, p, l, h, num(r[10]), "smae",
                        {"k_mg": r1(num(r[15])) or None, "na": r1(num(r[20])) or None}))
    return out, omit

def europeos(existentes):
    vistos = {norm(f["n"]) for f in existentes}
    out, dup = [], []
    for e in csv.DictReader(open(ROOT / "scripts" / "europeos.csv", encoding="utf-8"), delimiter=";"):
        if norm(e["nombre"]) in vistos: dup.append(e["nombre"]); continue
        sub = e["subgrupo"]; rk, rp, rl, rh = REF[sub]
        k, p, l, h, f = (float(e[x]) for x in ("kcal", "p", "l", "hc", "fibra"))
        # gramos de 1 equivalente: proteína en origen animal, grasa en grasas, energía en lo demás
        g = 100 * (rp / p if sub.startswith("A.O.A") else rl / l if sub.startswith("Grasas") else rk / k)
        g = round(g) if g < 20 else 5 * round(g / 5) if g < 100 else 10 * round(g / 10)
        gu = num(e["g_unidad"])
        if gu and e["unidad"] != "g":
            cant, unidad = g / gu, e["unidad"]
            cant = round(cant) if cant > 6 else round(cant * 2) / 2 if cant > 2 else round(cant * 4) / 4 or 0.25
            if e["unidad"] == "taza" and cant < 0.5: cant, unidad = g, "g"  # porciones muy chicas, mejor en gramos
        else:
            cant, unidad = g, "g"
        m = g / 100
        out.append(food(e["nombre"], sub, cant, unidad, g, k * m, p * m, l * m, h * m, f * m, "europa"))
    return out, dup

RENAL_COLS = {"energia (kcal)": "k", "proteina (g)": "p", "proteína (g)": "p", "lípidos (g)": "l", "hidratos de carbono (g)": "h",
              "fibra (g)": "f", "potasio (mg)": "k_mg", "sodio (mg)": "na", "fósforo (mg)": "fo", "fósforo mg": "fo", "agua (ml)": "agua",
              "agua ml": "agua", "agual ml": "agua", "calcio (mg)": "ca", "peso neto (g)": "pn"}

def renal(xlsm):
    wb = openpyxl.load_workbook(xlsm, read_only=True, data_only=True)
    rows = list(wb["Renal"].iter_rows(values_only=True))
    out, seccion, cols, malos = [], None, None, []
    for i, r in enumerate(rows[28:], start=28):
        vals = [c for c in r if c is not None]
        if not vals: continue
        if len(vals) == 1 and isinstance(vals[0], str) and vals[0].isupper():
            seccion, cols = re.sub(r"\s+", " ", vals[0]).strip().capitalize(), None; continue
        if str(r[0]).strip().upper() == "ALIMENTOS":
            cols = {}
            for j, h in enumerate(r):
                key = RENAL_COLS.get(re.sub(r"\s+", " ", str(h or "")).strip().lower())
                if key and key not in cols: cols[key] = j
            continue
        if not cols or not isinstance(r[0], str): continue
        v = {k: num(r[j]) for k, j in cols.items()}
        c = num(r[1])
        if c is None or v.get("k") is None or not v.get("pn"): malos.append((i, r[0], "sin datos")); continue
        # renglones corridos: macros imposibles para el peso
        if sum(x or 0 for x in (v.get("p"), v.get("l"), v.get("h"))) > v["pn"] * 1.05 or (v.get("k") or 0) > v["pn"] * 9.5:
            malos.append((i, r[0], "datos corridos")); continue
        if CHATARRA.search(r[0]): malos.append((i, r[0], "chatarra")); continue
        out.append(food(r[0], seccion, c, r[2] or "", v["pn"], v["k"], v.get("p"), v.get("l"), v.get("h"), v.get("f"), "renal",
                        {x: r1(v.get(x)) for x in ("k_mg", "na", "fo", "agua", "ca")}))
    return out, malos

if __name__ == "__main__":
    s, omit = smae(sys.argv[1], sys.argv[2])
    eu, dup = europeos(s)
    ids.clear()
    rn, malos = renal(sys.argv[1])
    OUT.mkdir(parents=True, exist_ok=True)
    quitados = []  # no saludables (criterio en scripts/saludable.py)
    def sanos(lista, sistema):
        q = [(sistema, f["g"], f["n"], motivo(f["n"], f["g"])) for f in lista if motivo(f["n"], f["g"])]
        quitados.extend(q)
        return [f for f in lista if not motivo(f["n"], f["g"])]
    s, eu, rn = sanos(s, "SMAE"), sanos(eu, "Europa"), sanos(rn, "Renal")
    if len(sys.argv) > 3:  # lista para que Cecy revise lo que se quitó
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Quitados"
        ws.append(["Sistema", "Subgrupo", "Alimento", "Motivo", "¿Regresar? (sí)"])
        for r in sorted(quitados, key=lambda r: (r[3], r[0], r[1], r[2])): ws.append(list(r))
        for col, w in zip("ABCDE", (9, 30, 55, 42, 14)): ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        wb.save(sys.argv[3])
    print("no saludables quitados:", len(quitados))
    # lac: tiene lactosa (para el filtro "Sin lactosa" de la paciente)
    LAC = re.compile(r"leche|yog|queso|cheddar|mozzarella|parmesano|gouda|emmental|halloumi|harzer|quark|skyr|k[eé]fir|jocoque|reques[oó]n|crema(?! de (cacahuate|almendra|cacao))|nata\b|mantequilla|butter|ghee|cajeta|lactium|activia|vitalin|whey|suero de leche|milch|lactobacil|shirota|jcoque|lerche", re.I)
    SIN = re.compile(r"deslact|sin lactosa|laktosefrei|soya|soja|almendra|avena|coco|arroz|anacardo|nuez de la india|vegan|vegetal|tofu|cacah|cacach|pera mantequilla", re.I)
    for f in s + eu + rn:
        if LAC.search(f["n"]) and not SIN.search(f["n"]): f["lac"] = 1
    for name, data in (("smae", s + eu), ("renal", rn)):
        (OUT / f"{name}.json").write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"SMAE: {len(s)} + europeos {len(eu)} (repetidos omitidos: {dup})")
    print("omitidos:", {k: (len(v) if isinstance(v, list) else v) for k, v in omit.items()})
    print("chatarra:", omit["chatarra"])
    print(f"Renal: {len(rn)}; omitidos {len(malos)}:", malos[:40])
